"""The literature review builder (Sprint 08).
The outline has one section for each sub-question. Each section lists the cards that carry its tag, grouped by the type of link,
with their checked quotes. The student writes the text. The AI writes nothing here.
A line that starts with "> " is a quote. It does not count as a word that the student wrote."""
import io
import json

from . import cards, cite, db, game

RELATION_LABEL = {"same_problem": "Same problem", "same_method": "Same method", "same_data": "Same data", "builds_on": "Builds on",
                  "compares_with": "Can be compared", "complements": "Complements", "other": "Other"}
GROUP_ORDER = list(RELATION_LABEL)
MAX_TEXT = 60000
AUTHOR_WORDS = 300  # the words that make a section count for the level "Author"


class WritingError(Exception):
    """Wrong input. The message is for the user."""


def count_words(text: str) -> int:
    """The words that the student wrote. Quote lines (they start with ">") are not counted."""
    return sum(len(line.split()) for line in (text or "").splitlines() if not line.lstrip().startswith(">"))


def quote_lines(text: str) -> list[str]:
    return [line.lstrip()[1:].strip() for line in (text or "").splitlines() if line.lstrip().startswith(">")]


def check_quotes(text: str) -> list[dict]:
    """Each quote line: is the quote in one of your PDFs? A quote that you changed by hand shows as not found."""
    pages = [cards.norm(t) for p in db.list_papers() for t in db.get_pages(p["id"])]
    out = []
    for line in quote_lines(text):
        body = line
        if body.endswith(")") and "(" in body:  # drop the citation at the end
            body = body[: body.rindex("(")]
        body = body.strip().strip("\"“”' ")
        n = cards.norm(body)
        out.append({"line": line, "found": bool(n) and len(n.split()) >= 3 and any(n in p for p in pages)})
    return out


# ------------------------------------------------------------------ the sources of a section
def _card_entries(pid: str, card_id: str) -> tuple[dict | None, str]:
    p = db.get_paper(pid)
    if not p:
        return None, ""
    card = db.get_card(pid) if not card_id else (db.get_extra(pid, card_id) or {}).get("card")
    return card, p["title"] + (f" ({card['focus']})" if card and card.get("focus") else "")


def sources(sub_question_id: str, style: str = "apa") -> list[dict]:
    """The cards with the tag of this sub-question, grouped by link type. Only quotes that the server verified."""
    tags = db.all_tags()
    papers = {p["id"]: p for p in db.list_papers()}
    order = {p["id"]: i + 1 for i, p in enumerate(sorted(papers.values(), key=lambda x: x["created_at"]))}
    mine = [(pid, cid) for pid, by in tags.items() for cid, ids in by.items() if sub_question_id in ids and pid in papers]
    relation: dict[str, str] = {}
    for pid, _ in mine:
        best = None
        for other, _ in mine:
            if other == pid:
                continue
            link = db.get_setting("link:" + ":".join(sorted([pid, other])))
            if link:
                rel = json.loads(link).get("relation", "other")
                if best is None or GROUP_ORDER.index(rel if rel in GROUP_ORDER else "other") < GROUP_ORDER.index(best):
                    best = rel if rel in GROUP_ORDER else "other"
        relation[pid] = best or "other"
    groups: dict[str, list[dict]] = {}
    for pid, cid in mine:
        card, title = _card_entries(pid, cid)
        if not card:
            continue
        quotes = []
        for name in cards.PAPER_FIELDS:
            f = (card.get("fields") or {}).get(name) or {}
            for e in f.get("evidence", []):
                if e.get("verified"):
                    c = cite.cite(papers[pid], e["page"], style, order[pid])
                    quotes.append({"field": name, "label": cards.FIELD_LABELS[name], "quote": e["quote"], "page": e["page"], "cite": c["text"], "complete": c["complete"]})
        groups.setdefault(relation[pid], []).append({"paper_id": pid, "card_id": cid, "title": title, "quotes": quotes})
    return [{"relation": r, "label": RELATION_LABEL[r], "cards": groups[r]} for r in GROUP_ORDER if r in groups]


# ------------------------------------------------------------------ the document
def outline() -> list[dict]:
    """One section for each sub-question. A section that exists keeps its text. No AI text: only headings."""
    doc = db.review_doc()
    have = {s["sub_question_id"] for s in db.review_sections(doc["id"])}
    subs = db.list_sub_questions()
    if not subs and not db.review_sections(doc["id"]):
        db.review_section_add(doc["id"], "Literature review", "")
    for s in subs:
        if s["id"] not in have:
            db.review_section_add(doc["id"], s["text"], s["id"])
    return db.review_sections(doc["id"])


def view(style: str = "apa") -> dict:
    doc = db.review_doc()
    secs = []
    for s in db.review_sections(doc["id"]):
        secs.append({**s, "words": count_words(s["text"]), "sources": sources(s["sub_question_id"], style) if s["sub_question_id"] else [],
                     "unverified_quotes": [q["line"] for q in check_quotes(s["text"]) if not q["found"]]})
    return {"doc": doc, "style": style, "sections": secs, "words": sum(s["words"] for s in secs)}


def save_text(sid: str, text: str | None, heading: str | None) -> dict:
    s = db.review_section_get(sid)
    if not s:
        raise WritingError("Section not found.")
    if text is not None:
        if len(text) > MAX_TEXT:
            raise WritingError(f"The text is longer than {MAX_TEXT} characters.")
        delta = count_words(text) - count_words(s["text"])
        db.review_section_update(sid, text=text)
        if delta > 0:  # the words of today, for the goal "words"
            db.writing_log_add(game.local_date(db.now()), sid, delta)
    if heading is not None:
        heading = " ".join(heading.split())
        if not heading or len(heading) > 200:
            raise WritingError("The heading must have 1 to 200 characters.")
        db.review_section_update(sid, heading=heading)
    return db.review_section_get(sid)


# ------------------------------------------------------------------ export
def _cited_papers(secs: list[dict]) -> list[dict]:
    titles = {p["id"]: p for p in db.list_papers()}
    ids = []
    for s in secs:
        for g in sources(s["sub_question_id"]) if s["sub_question_id"] else []:
            for c in g["cards"]:
                if c["paper_id"] not in ids:
                    ids.append(c["paper_id"])
    return [titles[i] for i in ids if i in titles]


def export_markdown(style: str = "apa") -> str:
    doc = db.review_doc()
    secs = db.review_sections(doc["id"])
    out = [f"# {doc['title']}", ""]
    for s in secs:
        out += [f"## {s['heading']}", "", (s["text"] or "").strip(), ""]
    refs = _cited_papers(secs)
    if refs:
        out += ["## References", ""]
        order = {p["id"]: i + 1 for i, p in enumerate(sorted(db.list_papers(), key=lambda x: x["created_at"]))}
        out += [f"- {cite.reference(p, style, order[p['id']])}" for p in sorted(refs, key=lambda p: (cite.surname((cite.authors_of(p) or ['~'])[0]).lower(), p.get("year") or ""))]
        out.append("")
    return "\n".join(out)


def export_docx(style: str = "apa") -> bytes:
    from docx import Document

    d = Document()
    doc = db.review_doc()
    d.add_heading(doc["title"], 0)
    secs = db.review_sections(doc["id"])
    for s in secs:
        d.add_heading(s["heading"], 1)
        for block in (s["text"] or "").split("\n"):
            if not block.strip():
                continue
            if block.lstrip().startswith(">"):
                d.add_paragraph(block.lstrip()[1:].strip(), style="Intense Quote")
            else:
                d.add_paragraph(block)
    refs = _cited_papers(secs)
    if refs:
        d.add_heading("References", 1)
        order = {p["id"]: i + 1 for i, p in enumerate(sorted(db.list_papers(), key=lambda x: x["created_at"]))}
        for p in refs:
            d.add_paragraph(cite.reference(p, style, order[p["id"]]))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
